"""processing_jobs / worker / PDF & DOCX evidence parsing / retry.

Main Spec §17 Phase 2: Job table and worker, PDF/DOCX parsing, retry.
Exercises the real dispatch in app/services/jobs.py end to end through the
HTTP API, the same way tests/test_smoke.py exercises the XLSX/plain-text
path.
"""

from __future__ import annotations

import io

import fitz
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.models import ProcessingJob


def _build_questionnaire_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Questionnaire"
    ws.append(["external_question_id", "question_text", "section", "is_required"])
    for row in rows:
        ws.append(
            [
                row.get("external_question_id", ""),
                row["question_text"],
                row.get("section", ""),
                row.get("is_required", True),
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_pdf(pages: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    buf = doc.tobytes()
    doc.close()
    return buf


def _build_docx(structure: list[tuple[str, str | None]]) -> bytes:
    doc = DocxDocument()
    for text, style in structure:
        doc.add_paragraph(text, style=style)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _create_case_with_electricity_question(client) -> tuple[str, str]:
    resp = client.post("/api/v1/cases", json={"title": "Case"})
    case_id = resp.json()["id"]

    xlsx_bytes = _build_questionnaire_xlsx(
        [{"question_text": "Report annual electricity consumption.", "is_required": True}]
    )
    resp = client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={
            "file": (
                "questionnaire.xlsx",
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"document_type": "QUESTIONNAIRE"},
    )
    question_id = client.get(f"/api/v1/cases/{case_id}/questions").json()[0]["id"]
    return case_id, question_id


# --------------------------------------------------------------------------- #
# processing_jobs created for every upload
# --------------------------------------------------------------------------- #


def test_upload_creates_a_processing_job_row(client, db_session):
    resp = client.post("/api/v1/cases", json={"title": "Case"})
    case_id = resp.json()["id"]

    xlsx_bytes = _build_questionnaire_xlsx(
        [{"question_text": "What is the total headcount?", "is_required": True}]
    )
    resp = client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={
            "file": (
                "q.xlsx",
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"document_type": "QUESTIONNAIRE"},
    )
    document = resp.json()
    assert document["latest_job_id"]

    job = db_session.get(ProcessingJob, document["latest_job_id"])
    assert job is not None
    assert job.job_type == "DOCUMENT_PARSE"
    assert job.status == "SUCCEEDED"
    assert job.started_at is not None
    assert job.finished_at is not None


# --------------------------------------------------------------------------- #
# PDF evidence
# --------------------------------------------------------------------------- #


def test_pdf_evidence_upload_creates_chunks_with_page_number(client, db_session):
    case_id, question_id = _create_case_with_electricity_question(client)

    pdf_bytes = _build_pdf(
        [
            "Employee headcount summary for FY2025.",
            "Total electricity consumption: 12,840 kWh at the Selangor site in January 2025.",
        ]
    )
    resp = client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={"file": ("utility-bill.pdf", pdf_bytes, "application/pdf")},
        data={"document_type": "UTILITY_BILL"},
    )
    assert resp.status_code == 201, resp.text
    document = resp.json()
    assert document["processing_status"] == "INDEXED"

    from app.models import Document, DocumentChunk

    doc_row = db_session.get(Document, document["id"])
    chunks = (
        db_session.query(DocumentChunk)
        .filter(DocumentChunk.document_id == doc_row.id)
        .order_by(DocumentChunk.sequence_no)
        .all()
    )
    assert len(chunks) == 2
    assert chunks[0].page_number == 1
    assert chunks[1].page_number == 2
    assert "electricity" in chunks[1].text.lower()

    resp = client.get(f"/api/v1/cases/{case_id}/questions")
    electricity_q = next(q for q in resp.json() if q["id"] == question_id)
    assert electricity_q["evidence_status"] == "PARTIAL"
    assert electricity_q["evidence_location"]["type"] == "page"
    assert electricity_q["evidence_location"]["page_number"] == 2


# --------------------------------------------------------------------------- #
# DOCX evidence
# --------------------------------------------------------------------------- #


def test_docx_evidence_upload_creates_chunks_with_heading_path(client, db_session):
    case_id, question_id = _create_case_with_electricity_question(client)

    docx_bytes = _build_docx(
        [
            ("Environmental Disclosures", "Heading 1"),
            (
                "Total electricity consumption: 12,840 kWh at the Selangor site "
                "in January 2025.",
                None,
            ),
        ]
    )
    resp = client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={
            "file": (
                "utility-bill.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"document_type": "UTILITY_BILL"},
    )
    assert resp.status_code == 201, resp.text
    document = resp.json()
    assert document["processing_status"] == "INDEXED"

    from app.models import Document, DocumentChunk

    doc_row = db_session.get(Document, document["id"])
    chunks = (
        db_session.query(DocumentChunk).filter(DocumentChunk.document_id == doc_row.id).all()
    )
    assert len(chunks) == 1
    assert chunks[0].page_number is None
    import json as _json

    assert _json.loads(chunks[0].heading_path) == ["Environmental Disclosures"]

    resp = client.get(f"/api/v1/cases/{case_id}/questions")
    electricity_q = next(q for q in resp.json() if q["id"] == question_id)
    assert electricity_q["evidence_status"] == "PARTIAL"
    assert electricity_q["evidence_location"]["type"] == "paragraph"
    assert electricity_q["evidence_location"]["heading_path"] == ["Environmental Disclosures"]


# --------------------------------------------------------------------------- #
# Parser failure -> FAILED job + NEEDS_MANUAL_REVIEW document, no crash
# --------------------------------------------------------------------------- #


def test_unparseable_pdf_marks_document_failed_without_crashing(client, db_session):
    resp = client.post("/api/v1/cases", json={"title": "Case"})
    case_id = resp.json()["id"]

    resp = client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={"file": ("bad.pdf", b"not a real pdf file", "application/pdf")},
        data={"document_type": "UTILITY_BILL"},
    )
    assert resp.status_code == 201, resp.text  # the request itself must not 500
    document = resp.json()
    assert document["processing_status"] == "NEEDS_MANUAL_REVIEW"
    assert document["error"]

    job = db_session.get(ProcessingJob, document["latest_job_id"])
    assert job.status == "FAILED"
    assert job.error_code == "DOCUMENT_PARSE_FAILED"
    assert job.error_message


# --------------------------------------------------------------------------- #
# Retry
# --------------------------------------------------------------------------- #


def test_retry_reattempts_processing_after_fixing_the_file(client, db_session):
    resp = client.post("/api/v1/cases", json={"title": "Case"})
    case_id = resp.json()["id"]

    resp = client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={"file": ("bad.pdf", b"not a real pdf file", "application/pdf")},
        data={"document_type": "UTILITY_BILL"},
    )
    document_id = resp.json()["id"]
    assert resp.json()["processing_status"] == "NEEDS_MANUAL_REVIEW"
    first_job_id = resp.json()["latest_job_id"]

    # Fix the stored file in place (simulating a corrected re-upload without
    # a new Document row — the retry endpoint re-parses the *stored* bytes).
    from app.models import Document
    from app.services import storage

    doc_row = db_session.get(Document, document_id)
    good_pdf = _build_pdf(["Total electricity consumption: 500 kWh."])
    storage.save(doc_row.storage_key, good_pdf)

    resp = client.post(f"/api/v1/cases/{case_id}/documents/{document_id}/retry")
    assert resp.status_code == 200, resp.text
    retried = resp.json()
    assert retried["processing_status"] == "INDEXED"
    assert retried["error"] is None
    assert retried["latest_job_id"] != first_job_id  # a new job attempt, not a mutated one

    # The failed job stays on record as history.
    first_job = db_session.get(ProcessingJob, first_job_id)
    assert first_job.status == "FAILED"

    new_job = db_session.get(ProcessingJob, retried["latest_job_id"])
    assert new_job.status == "SUCCEEDED"
    assert new_job.job_type == "DOCUMENT_INDEX"


def test_retry_rejects_a_document_that_is_not_in_a_retryable_state(client):
    resp = client.post("/api/v1/cases", json={"title": "Case"})
    case_id = resp.json()["id"]

    resp = client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={"file": ("note.txt", b"Some evidence text.", "text/plain")},
        data={"document_type": "OTHER"},
    )
    document_id = resp.json()["id"]
    assert resp.json()["processing_status"] == "INDEXED"

    resp = client.post(f"/api/v1/cases/{case_id}/documents/{document_id}/retry")
    assert resp.status_code == 409
    assert resp.json()["detail"]["error"]["code"] == "DOCUMENT_NOT_RETRYABLE"


def test_retry_on_unknown_document_returns_404(client):
    resp = client.post("/api/v1/cases", json={"title": "Case"})
    case_id = resp.json()["id"]

    resp = client.post(f"/api/v1/cases/{case_id}/documents/does-not-exist/retry")
    assert resp.status_code == 404
    assert resp.json()["detail"]["error"]["code"] == "DOCUMENT_NOT_FOUND"


# --------------------------------------------------------------------------- #
# GET /documents listing (used by the retry UI)
# --------------------------------------------------------------------------- #


def test_list_documents_returns_all_documents_for_a_case(client):
    resp = client.post("/api/v1/cases", json={"title": "Case"})
    case_id = resp.json()["id"]

    client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={"file": ("a.txt", b"Some text.", "text/plain")},
        data={"document_type": "OTHER"},
    )
    client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={"file": ("b.txt", b"Other text.", "text/plain")},
        data={"document_type": "OTHER"},
    )

    resp = client.get(f"/api/v1/cases/{case_id}/documents")
    assert resp.status_code == 200
    names = {doc["original_filename"] for doc in resp.json()}
    assert names == {"a.txt", "b.txt"}


# --------------------------------------------------------------------------- #
# worker.py claim_next_job (SQLite fallback path)
# --------------------------------------------------------------------------- #


def test_claim_next_job_claims_queued_job_and_marks_it_running(client, db_session):
    from app.services import jobs

    resp = client.post("/api/v1/cases", json={"title": "Case"})
    case_id = resp.json()["id"]

    job = jobs.create_job(db_session, case_id=case_id, job_type="DOCUMENT_INDEX")
    db_session.commit()
    assert job.status == "QUEUED"

    claimed = jobs.claim_next_job(db_session)
    assert claimed is not None
    assert claimed.id == job.id
    assert claimed.status == "RUNNING"
    assert claimed.attempt_count == 1

    # A second claim finds nothing else queued.
    assert jobs.claim_next_job(db_session) is None
